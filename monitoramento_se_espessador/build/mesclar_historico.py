"""Junta o histórico da planilha antiga com o da planilha nova (já usada) em historico.json.

Uso:  python3 mesclar_historico.py "Monitoramento SE - Após fator ESP.xlsx" "Monitoramento_SE_Espessador_Batch.xlsm"

Regras:
  - dias que estão só numa das planilhas entram como estão;
  - dia nas duas: vale o valor buscado no MES / lançado na planilha nova; campo vazio na nova é completado pela antiga;
  - quando os dois têm valor e eles diferem, fica o da nova e a coluna Origem registra o da planilha antiga.
"""
import os
import sys
import json
import datetime as dtm

import openpyxl

from extrair_historico import ABA, COLS, numero, lama

AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(AQUI, 'historico.json')
CAMPOS = ['prodC3', 'prodC2', 'espC3', 'espC2', 'eb4', 'eb6', 'tk01', 'tk02', 'fator', 'lamaC3', 'lamaC2', 'concC3',
          'concC2']
ROTULO = {'prodC3': 'produção C3', 'prodC2': 'produção C2', 'espC3': 'ESP C3', 'espC2': 'ESP C2', 'eb4': 'EB4',
          'eb6': 'EB6', 'tk01': 'TK 01', 'tk02': 'TK 02', 'fator': 'fator', 'lamaC3': 'lama C3', 'lamaC2': 'lama C2',
          'concC3': 'conc. C3', 'concC2': 'conc. C2'}
# aba Histórico da planilha nova: cabeçalho na linha 6, dados a partir da 7, coluna B em diante
H_TITULOS = {'Data': 'data', 'Produção C3 (t)': 'prodC3', 'Produção C2 (t)': 'prodC2', 'S.E. ESP C3': 'espC3',
             'S.E. ESP C2': 'espC2', 'S.E. batch EB4': 'eb4', 'S.E. batch EB6': 'eb6', 'TK 01 - EB6': 'tk01',
             'TK 02 - EB6': 'tk02', 'Fator': 'fator', 'Lama incorp. C3 (tms)': 'lamaC3',
             'Lama incorp. C2 (tms)': 'lamaC2', 'Concentrado C3 (tms)': 'concC3', 'Concentrado C2 (tms)': 'concC2',
             'Origem': 'origem', 'Atualizado em': 'atual', 'Observação': 'obs'}


def ler_antiga(caminho):
    ws = openpyxl.load_workbook(caminho)[ABA]
    dias = {}
    for r in range(2, ws.max_row + 1):
        d = ws['A%d' % r].value
        if not isinstance(d, dtm.datetime):
            continue
        reg = {k: numero(ws['%s%d' % (c, r)].value) for c, k in COLS.items()}
        if all(reg[k] is None for k in ('prodC3', 'prodC2', 'espC3', 'espC2', 'eb4', 'eb6')):
            continue
        reg['lamaC3'], reg['lamaC2'] = lama(ws['O%d' % r].value)
        dias[d.date().isoformat()] = {k: v for k, v in reg.items() if v is not None}
    return dias


def ler_nova(caminho):
    ws = openpyxl.load_workbook(caminho, data_only=True)['Histórico']
    cols = {}
    for c in range(1, ws.max_column + 1):
        t = ws.cell(6, c).value
        if t in H_TITULOS:
            cols[H_TITULOS[t]] = c
    dias = {}
    for r in range(7, ws.max_row + 1):
        d = ws.cell(r, cols['data']).value
        if not isinstance(d, dtm.datetime):
            continue
        reg = {}
        for k, c in cols.items():
            v = ws.cell(r, c).value
            if k in CAMPOS:
                v = numero(v)
            if v not in (None, ''):
                reg[k] = v
        reg.pop('data', None)
        dias[d.date().isoformat()] = reg
    return dias


def main():
    if len(sys.argv) < 3:
        raise SystemExit(__doc__)
    antiga, nova = ler_antiga(sys.argv[1]), ler_nova(sys.argv[2])
    saida, notas = [], []
    for dia in sorted(set(antiga) | set(nova)):
        a, n = antiga.get(dia), nova.get(dia)
        if n is None:
            reg = dict(a, origem='Planilha anterior')
        elif a is None or n.get('origem') == 'Importado':
            reg = dict(n)
        else:
            reg = dict(n)
            dif = []
            for k in CAMPOS:
                if k not in reg and k in a:
                    reg[k] = a[k]
                    dif.append('%s da planilha anterior (MES sem valor)' % ROTULO[k])
                elif k in reg and k in a and abs(reg[k] - a[k]) > 0.5:
                    dif.append('%s na planilha anterior: %s' % (ROTULO[k], ('%.2f' % a[k]).rstrip('0').rstrip('.')))
            if dif:
                # fica na coluna Origem (a Observacao vai para os comentarios do painel)
                reg['origem'] = '%s + planilha anterior (%s)' % (reg.get('origem', 'MES'), '; '.join(dif))
                notas.append('%s: %s' % (dia, '; '.join(dif)))
        if isinstance(reg.get('atual'), dtm.datetime):
            reg['atual'] = reg['atual'].isoformat(timespec='minutes')
        reg['data'] = dia
        saida.append(reg)
    with open(SAIDA, 'w', encoding='utf-8') as fh:
        json.dump(saida, fh, ensure_ascii=False, indent=0)
    so_a = sorted(set(antiga) - set(nova))
    print('%d dias (%s a %s) -> %s' % (len(saida), saida[0]['data'], saida[-1]['data'], SAIDA))
    print('só na planilha anterior: %d (%s a %s)' % (len(so_a), so_a[0], so_a[-1]) if so_a else 'só na anterior: 0')
    for t in notas:
        print('  ' + t)


if __name__ == '__main__':
    main()
