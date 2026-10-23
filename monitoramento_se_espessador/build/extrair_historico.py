"""Importa a base da planilha atual (aba "Média ponderada") para historico.json.

Uso:  python3 extrair_historico.py "Monitoramento SE - Após fator ESP 6.2.xlsx"

Colunas da planilha antiga -> campos do historico:
  B Produção C3 · C Produção C2 · D Média ESP C3 · E Média ESP C2 · G Média batch EB4 · I Média batch EB6
  L TK 01 · M TK 02 · N Fator · O Lama incorporada (=C3+C2: a 1a parcela vai para C3, a 2a para C2)
  P Concentrado C3 · Q Concentrado C2
F, H, J, K e R são calculadas e ficam de fora (a planilha nova calcula de novo).
"""
import os
import re
import sys
import json
import datetime as dtm

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
SAIDA = os.path.join(AQUI, 'historico.json')
ABA = 'Média ponderada'
COLS = {'B': 'prodC3', 'C': 'prodC2', 'D': 'espC3', 'E': 'espC2', 'G': 'eb4', 'I': 'eb6',
        'L': 'tk01', 'M': 'tk02', 'N': 'fator', 'P': 'concC3', 'Q': 'concC2'}


def numero(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    return None          # '-' e textos: sem valor


def lama(v):
    """O: número (só C3) ou fórmula de soma '=412+50' (C3 + C2)."""
    if isinstance(v, str) and v.startswith('='):
        partes = [float(p) for p in re.findall(r'[-+]?\d+(?:\.\d+)?', v[1:])]
        c3 = partes[0] if partes else None
        c2 = sum(partes[1:]) if len(partes) > 1 else None
        return c3, c2
    return numero(v), None


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    ws = openpyxl.load_workbook(sys.argv[1])[ABA]
    dias = []
    for r in range(2, ws.max_row + 1):
        d = ws['A%d' % r].value
        if not isinstance(d, dtm.datetime):
            continue
        reg = {k: numero(ws['%s%d' % (c, r)].value) for c, k in COLS.items()}
        if all(reg[k] is None for k in ('prodC3', 'prodC2', 'espC3', 'espC2', 'eb4', 'eb6')):
            continue
        reg['lamaC3'], reg['lamaC2'] = lama(ws['O%d' % r].value)
        reg = {k: v for k, v in reg.items() if v is not None}
        reg['data'] = d.date().isoformat()
        dias.append(reg)
    with open(SAIDA, 'w', encoding='utf-8') as fh:
        json.dump(dias, fh, ensure_ascii=False, indent=0)
    print('%d dias (%s a %s) -> %s' % (len(dias), dias[0]['data'], dias[-1]['data'], SAIDA))


if __name__ == '__main__':
    main()
