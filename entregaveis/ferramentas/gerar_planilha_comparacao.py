"""Gera a planilha de comparação entre o módulo de compressão automática e o resultado do MES.

Uso: python gerar_planilha_comparacao.py [arquivo_saida.xlsx]

A planilha sai preenchida com DADOS SIMULADOS para demonstrar o funcionamento.
Para usar com dados reais, cole as exportações do PIMS nas abas Modulo, MES e Amostragem.
"""
import math
import random
import sys
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = sys.argv[1] if len(sys.argv) > 1 else "comparacao_modulo_mes.xlsx"
N = 5001          # última linha aceita nas abas de dados
NJ = 501          # última linha da aba Janelas

F = "Arial"
CAB = PatternFill("solid", fgColor="1F2630")
ENTRADA = PatternFill("solid", fgColor="FFF2CC")
CALC = PatternFill("solid", fgColor="F2F2F2")
FINO = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
DH = "dd/mm/yyyy hh:mm"

wb = Workbook()


def cabecalho(ws, titulos, larguras):
    for c, (t, w) in enumerate(zip(titulos, larguras), 1):
        cel = ws.cell(row=1, column=c, value=t)
        cel.font = Font(name=F, bold=True, color="FFFFFF")
        cel.fill = CAB
        cel.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cel.column_letter].width = w
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "A2"


def fonte(ws):
    for linha in ws.iter_rows():
        for c in linha:
            if c.font is None or c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold if c.font else False, color=c.font.color if c.font else None,
                              size=c.font.size if c.font else 11, italic=c.font.italic if c.font else False)


# ---------------------------------------------------------------- Leia-me
ws = wb.active
ws.title = "Leia-me"
texto = [
    ("Comparação módulo de compressão automática x laboratório (MES), PQ Usina 4", True),
    ("", False),
    ("ATENÇÃO: os dados das abas Modulo, MES e Amostragem são SIMULADOS, apenas para demonstrar a planilha.", True),
    ("", False),
    ("Como usar", True),
    ("1. Exporte do PIMS (Aspen Process Data no Excel ou exportação do Process Explorer) o mesmo período para as tags abaixo.", False),
    ("2. Apague os dados simulados e cole os dados reais nas células amarelas das abas Modulo, MES e Amostragem, a partir da linha 2.", False),
    ("3. Ajuste os parâmetros na aba Parametros (células amarelas). Os valores marcados como 'a confirmar' são hipóteses.", False),
    ("4. A aba Janelas monta uma linha por resultado do MES, calcula a média do módulo na mesma janela e decide se a janela é válida.", False),
    ("5. A aba Resumo traz o viés, os limites de concordância (Bland-Altman), o teste t pareado e os gráficos.", False),
    ("", False),
    ("Tags de origem", True),
    ("Modulo: 307-GI-001 (carga média, daN/pelota), 307-GI-002 (desvio padrão), qualidade de 307-GI-001. Um ensaio a cada 30 min.", False),
    ("MES: M4710050020-0031-HHLFU (Compressão, mapa IP_MESVALOR, servidor UBU). Um resultado a cada 2 h.", False),
    ("Amostragem: modo de U04-07CR001 e U04-07CR002 (mapa IP_PIDMAP03: Automático ou Manual), falha do sistema (U04-07CR001, U04-07CR002, U04-07AL001, U04-07TP018 pelo mapa IP_PIDMAP01) e alimentação 406-FX-001 (t/h).", False),
    ("Sugestão de amostragem na exportação da aba Amostragem: um valor a cada 10 min (interpolado ou último valor).", False),
    ("", False),
    ("Legenda", True),
    ("Células amarelas: entrada de dados ou parâmetros. Células cinza: fórmulas, não editar.", False),
    ("Janela válida: resultado MES presente, módulo com o mínimo de ensaios Good, amostragem em automático e sem falha durante a janela e usina com alimentação acima do mínimo.", False),
]
for i, (t, b) in enumerate(texto, 1):
    ws.cell(row=i, column=1, value=t).font = Font(name=F, bold=b, size=13 if i == 1 else 11,
                                                   color="C00000" if t.startswith("ATENÇÃO") else None)
ws.column_dimensions["A"].width = 150

# ---------------------------------------------------------------- Parametros
wp = wb.create_sheet("Parametros")
cabecalho(wp, ["Parâmetro", "Valor", "Unidade", "Observação"], [62, 14, 14, 70])
params = [
    ("Duração da janela do MES", 2, "h", "Um resultado do MES a cada 2 h (informado pelo LCE)"),
    ("Início da janela em relação ao horário do resultado MES", -2, "h", "A confirmar. -2 significa que o resultado das 10:00 representa a amostra composta de 08:00 a 10:00. Use 0 se o horário marcar o início."),
    ("Mínimo de ensaios do módulo na janela", 3, "ensaios", "Esperados 4 ensaios (um a cada 30 min) por janela de 2 h"),
    ("Fração mínima do tempo em automático e sem falha", 1, "", "1 = 100 % da janela em automático e sem falha"),
    ("Alimentação mínima (406-FX-001)", 300, "t/h", "A confirmar. Abaixo disso a usina é considerada parada ou em baixa carga"),
    ("Qualidade aceita do módulo", "Good", "", "Ensaios com qualidade diferente são descartados"),
    ("Pelotas por ensaio no laboratório", 60, "pelotas", "A confirmar com o procedimento do LCE"),
    ("Pelotas por ensaio no módulo", 20, "pelotas", "A confirmar com o fornecedor do módulo"),
]
for r, (n, v, u, o) in enumerate(params, 2):
    wp.cell(row=r, column=1, value=n)
    c = wp.cell(row=r, column=2, value=v)
    c.fill = ENTRADA
    c.border = BORDA
    wp.cell(row=r, column=3, value=u)
    wp.cell(row=r, column=4, value=o).alignment = Alignment(wrap_text=True)
wp["B5"].number_format = "0%"
P_DUR, P_INI, P_MIN, P_AUTO, P_ALIM, P_QUAL, P_NLAB, P_NMOD = ("Parametros!$B$%d" % i for i in range(2, 10))

# ---------------------------------------------------------------- dados simulados
random.seed(42)
t0 = datetime(2026, 10, 20, 0, 0)
dias = 4


def qualidade_real(t):
    h = (t - t0).total_seconds() / 3600
    return 285 + 18 * math.sin(h / 7) + 8 * math.sin(h / 2.3)


manual = (datetime(2026, 10, 21, 8), datetime(2026, 10, 21, 14))
falha = (datetime(2026, 10, 22, 20), datetime(2026, 10, 22, 22, 30))
parada = (datetime(2026, 10, 23, 2), datetime(2026, 10, 23, 6))
dentro = lambda t, p: p[0] <= t < p[1]

mod_rows, mes_rows, am_rows = [], [], []
t = t0 + timedelta(minutes=30)
while t <= t0 + timedelta(days=dias):
    if not dentro(t, parada):
        q = "Bad" if random.random() < 0.03 else "Good"
        mod_rows.append((t, round(qualidade_real(t) + 8 + random.gauss(0, 9), 1), round(random.gauss(52, 5), 1), q))
    t += timedelta(minutes=30)
t = t0 + timedelta(hours=2)
while t <= t0 + timedelta(days=dias):
    if not dentro(t - timedelta(minutes=1), parada):
        base = sum(qualidade_real(t - timedelta(minutes=m)) for m in range(0, 120, 10)) / 12
        mes_rows.append((t, round(base + random.gauss(0, 9))))
    t += timedelta(hours=2)
t = t0
while t < t0 + timedelta(days=dias):
    m2 = "Manual" if dentro(t, manual) else "Automático"
    fl = 1 if dentro(t, falha) else 0
    al = 0 if dentro(t, parada) else round(1190 + random.gauss(0, 25), 1)
    am_rows.append((t, "Automático", m2, fl, al))
    t += timedelta(minutes=10)

# ---------------------------------------------------------------- Modulo
wm = wb.create_sheet("Modulo")
cabecalho(wm, ["Data e hora do ensaio", "Carga média 307-GI-001 (daN/pelota)", "Desvio padrão 307-GI-002 (daN/pelota)", "Qualidade 307-GI-001"], [20, 18, 18, 14])
for r, row in enumerate(mod_rows, 2):
    for c, v in enumerate(row, 1):
        cel = wm.cell(row=r, column=c, value=v)
        cel.fill = ENTRADA
    wm.cell(row=r, column=1).number_format = DH
    wm.cell(row=r, column=2).number_format = "0.0"
    wm.cell(row=r, column=3).number_format = "0.0"
wm["A1"].comment = Comment("Dados simulados. Substituir pela exportação do PIMS.", "LCE")

# ---------------------------------------------------------------- MES
we = wb.create_sheet("MES")
cabecalho(we, ["Data e hora do resultado", "Compressão MES M4710050020-0031-HHLFU (daN/pelota)"], [20, 22])
for r, row in enumerate(mes_rows, 2):
    for c, v in enumerate(row, 1):
        we.cell(row=r, column=c, value=v).fill = ENTRADA
    we.cell(row=r, column=1).number_format = DH
we["A1"].comment = Comment("Dados simulados. Um resultado por linha, em ordem cronológica, sem linhas em branco no meio.", "LCE")

# ---------------------------------------------------------------- Amostragem
wa = wb.create_sheet("Amostragem")
cabecalho(wa, ["Data e hora", "Modo 07CR001 (Automático/Manual)", "Modo 07CR002 (Automático/Manual)", "Falha no sistema (0 = não, 1 = sim)", "Alimentação 406-FX-001 (t/h)"], [20, 18, 18, 16, 16])
for r, row in enumerate(am_rows, 2):
    for c, v in enumerate(row, 1):
        wa.cell(row=r, column=c, value=v).fill = ENTRADA
    wa.cell(row=r, column=1).number_format = DH
    wa.cell(row=r, column=5).number_format = "0.0"
wa["D1"].comment = Comment("1 quando qualquer equipamento (07CR001, 07CR002, 07AL001, 07TP018) estiver em falha no instante.", "LCE")

# ---------------------------------------------------------------- Janelas
wj = wb.create_sheet("Janelas")
cab = ["Horário do resultado MES", "Compressão MES (daN/pelota)", "Início da janela", "Fim da janela",
       "Ensaios do módulo (Good)", "Média do módulo (daN/pelota)", "DP médio do módulo (daN/pelota)",
       "Fração em automático e sem falha", "Alimentação média (t/h)", "Situação da janela",
       "Diferença módulo menos MES (daN/pelota)", "Diferença relativa", "Média dos dois métodos (daN/pelota)",
       "Fora dos limites de concordância?",
       "X gráfico", "MES gráfico", "Módulo gráfico", "Diferença gráfico", "Média gráfico", "MES válida", "Módulo válido"]
cabecalho(wj, cab, [18, 13, 18, 18, 11, 13, 13, 13, 13, 40, 14, 11, 13, 13, 8, 8, 8, 8, 8, 8, 8])
MA, MB, MC, MD = ("Modulo!$%s$2:$%s$%d" % (c, c, N) for c in "ABCD")
AA, AB, AC, AD, AE = ("Amostragem!$%s$2:$%s$%d" % (c, c, N * 3) for c in "ABCDE")
for r in range(2, NJ + 1):
    f = {
        "A": f'=IF(MES!A{r}="","",MES!A{r})',
        "B": f'=IF(A{r}="","",MES!B{r})',
        "C": f'=IF(A{r}="","",A{r}+{P_INI}/24)',
        "D": f'=IF(A{r}="","",C{r}+{P_DUR}/24)',
        "E": f'=IF(A{r}="","",COUNTIFS({MA},">="&C{r},{MA},"<"&D{r},{MD},{P_QUAL}))',
        "F": f'=IF(A{r}="","",IF(E{r}=0,"",AVERAGEIFS({MB},{MA},">="&C{r},{MA},"<"&D{r},{MD},{P_QUAL})))',
        "G": f'=IF(A{r}="","",IF(E{r}=0,"",AVERAGEIFS({MC},{MA},">="&C{r},{MA},"<"&D{r},{MD},{P_QUAL})))',
        "H": f'=IF(A{r}="","",IFERROR(COUNTIFS({AA},">="&C{r},{AA},"<"&D{r},{AB},"Automático",{AC},"Automático",{AD},0)/COUNTIFS({AA},">="&C{r},{AA},"<"&D{r}),0))',
        "I": f'=IF(A{r}="","",IFERROR(AVERAGEIFS({AE},{AA},">="&C{r},{AA},"<"&D{r}),""))',
        "J": (f'=IF(A{r}="","",IF(B{r}="","Inválida: sem resultado MES",IF(E{r}<{P_MIN},"Inválida: poucos ensaios do módulo",'
              f'IF(H{r}<{P_AUTO},"Inválida: amostragem em manual ou em falha",IF(OR(I{r}="",I{r}<{P_ALIM}),"Inválida: usina parada ou baixa alimentação","Válida")))))'),
        "K": f'=IF(J{r}="Válida",F{r}-B{r},"")',
        "L": f'=IF(J{r}="Válida",K{r}/B{r},"")',
        "M": f'=IF(J{r}="Válida",(F{r}+B{r})/2,"")',
        "N": f'=IF(K{r}="","",IF(OR(K{r}>Resumo!$B$16,K{r}<Resumo!$B$15),"Sim","Não"))',
        "O": f'=IF(A{r}="",NA(),A{r})',
        "P": f'=IF(A{r}="",NA(),B{r})',
        "Q": f'=IF(F{r}="",NA(),F{r})',
        "R": f'=IF(K{r}="",NA(),K{r})',
        "S": f'=IF(M{r}="",NA(),M{r})',
        "T": f'=IF(J{r}="Válida",B{r},"")',
        "U": f'=IF(J{r}="Válida",F{r},"")',
    }
    for col, formula in f.items():
        c = wj[f"{col}{r}"]
        c.value = formula
        c.fill = CALC
    for col in "ACDO":
        wj[f"{col}{r}"].number_format = DH
    for col in "BFGIKMPQRSTU":
        wj[f"{col}{r}"].number_format = "0.0"
    wj[f"H{r}"].number_format = "0%"
    wj[f"L{r}"].number_format = "0.0%"
for col in "OPQRSTU":
    wj.column_dimensions[col].hidden = True

# ---------------------------------------------------------------- Resumo
wr = wb.create_sheet("Resumo", 1)
wr.column_dimensions["A"].width = 64
wr.column_dimensions["B"].width = 18
wr.column_dimensions["C"].width = 70
J = lambda c: f"Janelas!${c}$2:${c}${NJ}"
linhas = [
    ("Comparação módulo x MES: resumo do período", None, None),
    ("Período analisado: início", f"=IFERROR(MIN({J('A')}),\"\")", DH),
    ("Período analisado: fim", f"=IFERROR(MAX({J('A')}),\"\")", DH),
    ("Janelas com resultado MES", f"=COUNT({J('A')})", "0"),
    ("Janelas válidas", f'=COUNTIF({J("J")},"Válida")', "0"),
    ("Percentual de janelas válidas", "=IF(B4=0,\"\",B5/B4)", "0%"),
    ("Inválidas: amostragem em manual ou em falha", f'=COUNTIF({J("J")},"Inválida: amostragem*")', "0"),
    ("Inválidas: poucos ensaios do módulo", f'=COUNTIF({J("J")},"Inválida: poucos*")', "0"),
    ("Inválidas: usina parada ou baixa alimentação", f'=COUNTIF({J("J")},"Inválida: usina*")', "0"),
    ("Inválidas: sem resultado MES", f'=COUNTIF({J("J")},"Inválida: sem*")', "0"),
    ("", None, None),
    ("Viés médio, módulo menos MES (daN/pelota)", f"=IFERROR(AVERAGE({J('K')}),\"\")", "0.0"),
    ("Desvio padrão das diferenças (daN/pelota)", f"=IFERROR(STDEV({J('K')}),\"\")", "0.0"),
    ("Viés relativo médio", f"=IFERROR(AVERAGE({J('L')}),\"\")", "0.0%"),
    ("Limite inferior de concordância (viés menos 1,96 DP)", "=IFERROR(B12-1.96*B13,\"\")", "0.0"),
    ("Limite superior de concordância (viés mais 1,96 DP)", "=IFERROR(B12+1.96*B13,\"\")", "0.0"),
    ("Teste t pareado: estatística t", "=IFERROR(B12/(B13/SQRT(B5)),\"\")", "0.00"),
    ("Teste t pareado: p-valor (bilateral)", "=IFERROR(TDIST(ABS(B17),B5-1,2),\"\")", "0.0000"),
    ("Correlação entre módulo e MES (janelas válidas)", f"=IFERROR(CORREL({J("T")},{J("U")}),\"\")", "0.00"),
    ("Janelas válidas fora dos limites de concordância", f'=COUNTIF({J("N")},"Sim")', "0"),
    ("Diferença esperada só pelo acaso das pelotas, 95 % (daN/pelota)",
     f"=IFERROR(1.96*AVERAGE({J('G')})*SQRT(1/{P_NLAB}+1/({P_NMOD}*AVERAGE({J('E')}))),\"\")", "0.0"),
    ("Conclusão", '=IF(B5<10,"Poucas janelas válidas. Recomendado: no mínimo 30.",IF(B18<0.05,"Há viés estatisticamente significativo entre módulo e MES.","Não há evidência de viés entre módulo e MES."))&IF(B5<30," Amostra ainda pequena.","")', None),
]
notas = {
    12: "Positivo: o módulo mede mais que o laboratório.",
    13: "Inclui a variação natural entre as pelotas e a diferença de amostra entre os dois métodos.",
    15: "Faixa em que se esperam 95 % das diferenças, se o comportamento se mantiver.",
    17: "Compara o viés médio com zero.",
    18: "Abaixo de 0,05: o viés é estatisticamente significativo.",
    19: "Depende da faixa de variação do período: em período estável a correlação tende a ser baixa mesmo com boa concordância.",
    21: "Referência de quanto os dois métodos podem diferir mesmo medindo o mesmo material, só pelas pelotas sorteadas. Usa o DP médio do módulo e as quantidades de pelotas dos Parametros. Se o DP das diferenças (B13) for muito maior que isso, há outra fonte de diferença (amostra diferente, calibração ou horário).",
}
for r, (rot, form, fmt) in enumerate(linhas, 1):
    a = wr.cell(row=r, column=1, value=rot)
    a.font = Font(name=F, bold=(r == 1 or r == 22), size=13 if r == 1 else 11)
    if form is not None:
        b = wr.cell(row=r, column=2, value=form)
        b.fill = CALC
        b.border = BORDA
        if fmt:
            b.number_format = fmt
    if r in notas:
        wr.cell(row=r, column=3, value=notas[r]).alignment = Alignment(wrap_text=True, vertical="top")
wr["B22"].alignment = Alignment(wrap_text=False)


def grafico_disperso(titulo, x_tit, y_tit, x_col, ys, ancora, estilo_linha=False, eixo_data=False):
    ch = ScatterChart()
    ch.title = titulo
    ch.style = 13
    ch.x_axis.title = x_tit
    ch.y_axis.title = y_tit
    ch.height, ch.width = 9, 22
    xref = Reference(wj, min_col=x_col, min_row=2, max_row=NJ)
    for col, nome, cor, so_marcador in ys:
        yref = Reference(wj, min_col=col, min_row=2, max_row=NJ)
        s = Series(yref, xref, title=nome)
        s.marker.symbol = "circle"
        s.marker.size = 5
        s.marker.graphicalProperties.solidFill = cor
        s.marker.graphicalProperties.line.solidFill = cor
        if so_marcador:
            s.graphicalProperties.line.noFill = True
        else:
            s.graphicalProperties.line.solidFill = cor
        ch.series.append(s)
    if eixo_data:
        ch.x_axis.number_format = "dd/mm hh:mm"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    wr.add_chart(ch, ancora)


grafico_disperso("Compressão: MES x média do módulo na mesma janela", "Horário do resultado MES", "daN/pelota",
                 15, [(16, "MES", "C00000", False), (17, "Módulo (média da janela)", "1F77B4", False)], "A25", eixo_data=True)
grafico_disperso("Bland-Altman: diferença x média dos dois métodos", "Média dos dois métodos (daN/pelota)",
                 "Módulo menos MES (daN/pelota)", 19, [(18, "Janelas válidas", "1F77B4", True)], "A44")

dv = DataValidation(type="list", formula1='"Good"', allow_blank=False)
wp.add_data_validation(dv)
dv.add("B7")

for w in wb.worksheets:
    fonte(w)
wb.save(SAIDA)
print(f"Planilha gerada: {SAIDA} ({len(mod_rows)} ensaios do módulo, {len(mes_rows)} resultados MES, {len(am_rows)} registros de amostragem)")
