"""Motor de verificação: avalia cada fórmula da BD_Afericoes com a biblioteca `formulas`,
usando os valores gravados nas abas dos operadores. Gera cache.json (valores p/ <v>)."""
import sys, json, datetime, re
import numpy as np, openpyxl, formulas, schedula as sh
from openpyxl.utils import column_index_from_string as CI, get_column_letter as L
import spec

wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
E0 = datetime.datetime(1899, 12, 30)
def conv(v):
    if isinstance(v, datetime.datetime): d = v - E0; return d.days + d.seconds / 86400
    if isinstance(v, datetime.time): return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    if v is None: return sh.EMPTY
    return v
VAL = {}
for ws in wb.worksheets:
    if ws.title.startswith("BD_"): continue
    for row in ws.iter_rows():
        for c in row:
            if hasattr(c, "coordinate") and c.value is not None:
                VAL[(ws.title.upper(), c.coordinate)] = conv(c.value)
for i, t in enumerate(spec.LIM):              # BD_Limites: A..I
    for j, v in enumerate(t):
        if v is not None: VAL[("BD_LIMITES", "%s%d" % (L(j + 1), i + 2))] = v

rows = spec.resolve()
BD = {}
def get(sheet, ref):
    if sheet == "BD_AFERICOES": return BD.get(ref, sh.EMPTY)
    return VAL.get((sheet, ref), sh.EMPTY)
def resolve_input(name):
    if "!" in name:
        s, ref = name.rsplit("!", 1); s = s.strip("'")
    else:
        s, ref = "BD_AFERICOES", name
    if ":" in ref:
        a, b = ref.split(":")
        ca, ra = re.match(r"([A-Z]+)(\d+)", a).groups(); cb, rb = re.match(r"([A-Z]+)(\d+)", b).groups()
        return np.array([[get(s, "%s%d" % (L(c), r)) for c in range(CI(ca), CI(cb) + 1)]
                         for r in range(int(ra), int(rb) + 1)], dtype=object)
    return get(s, ref)
CACHE = {}
comp = {}
def scalar(x):
    while isinstance(x, np.ndarray): x = x.flat[0] if x.size else ""
    if x is sh.EMPTY: return ""
    if isinstance(x, (np.floating, np.integer)): x = x.item()
    if isinstance(x, float) and x == int(x) and abs(x) < 1e15: x = int(x)
    return x
ORDER = [c for c in "ABCDEFGHJKLNOPQ"] + ["I", "M"]
for n, row in enumerate(rows, start=2):
    for c in ORDER:
        v = row.get(c)
        ref = "%s%d" % (c, n)
        if isinstance(v, tuple):
            f = comp.get(v[1])
            if f is None:
                f = comp[v[1]] = formulas.Parser().ast("=" + v[1])[1].compile()
            out = scalar(f(*[resolve_input(i) for i in f.inputs]))
            if not isinstance(out, (int, float, str)):
                raise SystemExit("valor estranho %s %r (%s)" % (ref, out, v[1]))
            if isinstance(out, str) and "TRIM(" in v[1]:
                out = re.sub(" +", " ", out.strip(" "))         # TRIM do Excel também junta espaços internos
            if c == "E" and row["A"] == "Verificação Peneiradores":
                out = out.replace(".", ",")                      # número -> texto no Excel pt-BR usa vírgula
            BD[ref] = out
            CACHE[ref] = out
        elif v is not None:
            BD[ref] = v
json.dump({"BD_Afericoes": CACHE}, open(sys.argv[2], "w"), ensure_ascii=False)
erros = {k: v for k, v in CACHE.items() if isinstance(v, str) and v.startswith("#")}
print("fórmulas:", len(CACHE), "erros:", len(erros), list(erros.items())[:10])
