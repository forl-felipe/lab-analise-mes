# -*- coding: utf-8 -*-
"""Gera as planilhas mensais de calibração a partir da planilha de SETEMBRO (layout de referência).

    python3 mensal.py setembro BASE.xlsm SAIDA.xlsm
    python3 mensal.py padrao   BASE.xlsm SAIDA.xlsm
    python3 mensal.py outubro  BASE.xlsm SAIDA.xlsm OUTUBRO_ORIGINAL.xlsm

Edita o XML direto (sem salvar pelo openpyxl): VBA, gráficos, desenhos, validações e formatação ficam intactos.
1. abas dos operadores: limpa / copia lançamentos (só células de preenchimento, nunca fórmulas)
2. BD_Afericoes e BD_Limites: refeitas com spec.py (mesmas tabelas tbl_Afericoes e tbl_Limites)
3. valores gravados das fórmulas calculados por avalia.py (é o que o Power BI lê)
"""
import sys, os, re, io, json, zipfile, datetime, subprocess, copy
from xml.sax.saxutils import escape
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.utils.cell import coordinate_from_string as CF
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import spec
from limpeza import ENTRADAS, OUTUBRO

MODO, BASE, SAIDA = sys.argv[1], sys.argv[2], sys.argv[3]
FONTE_OUT = sys.argv[4] if len(sys.argv) > 4 else None
E0 = datetime.datetime(1899, 12, 30)
LOG = []

z = zipfile.ZipFile(BASE)
files = {n: z.read(n) for n in z.namelist()}
infos = {i.filename: i for i in z.infolist()}

# ── mapa aba -> arquivo ───────────────────────────────────────────────────
wbx = files["xl/workbook.xml"].decode("utf8")
rels = dict(re.findall(r'<Relationship Id="(rId\d+)"[^>]*?Target="([^"]+)"', files["xl/_rels/workbook.xml.rels"].decode("utf8")))
rels.update({a: b for b, a in re.findall(r'<Relationship Target="([^"]+)"[^>]*?Id="(rId\d+)"', files["xl/_rels/workbook.xml.rels"].decode("utf8"))})
ABA = {}
for nome, rid in re.findall(r'<sheet name="([^"]+)"[^>]*?r:id="(rId\d+)"', wbx):
    ABA[nome.replace("&amp;", "&")] = "xl/" + rels[rid]
assert "BD_Afericoes" in ABA and "BD_Limites" in ABA, ABA

# ── shared strings ────────────────────────────────────────────────────────
sst = files["xl/sharedStrings.xml"].decode("utf8")
SIS = re.findall(r"<si>.*?</si>|<si/>", sst, re.S)
IDX = {}
for i, si in enumerate(SIS):
    m = re.fullmatch(r'<si><t(?: xml:space="preserve")?>(.*?)</t></si>', si, re.S)
    if m: IDX.setdefault(m.group(1), i)
NOVOS = []; NREFS = [0]
def s_id(txt):
    NREFS[0] += 1
    k = escape(txt)
    if k not in IDX:
        IDX[k] = len(SIS) + len(NOVOS)
        sp = ' xml:space="preserve"' if txt != txt.strip() or "\n" in txt else ""
        NOVOS.append("<si><t%s>%s</t></si>" % (sp, k))
    return IDX[k]
def s_txt(i):
    si = SIS[i]
    return "".join(re.findall(r"<t[^>]*>(.*?)</t>", si, re.S)).replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")

# ── edição de células no XML das abas (cirurgia de texto: o resto do arquivo fica byte a byte igual) ──
TXT = {}
def xml(aba):
    p = ABA[aba]
    if p not in TXT: TXT[p] = files[p].decode("utf8")
    return p

def _acha(aba, ref):
    p = xml(aba); x = TXT[p]
    m = re.search(r'<c r="%s"(?=[ />])([^>]*?)(/>|>(.*?)</c>)' % ref, x, re.S)
    return p, m

def _attrs(a):
    return re.findall(r'([\w:]+)="([^"]*)"', a)

def celula_info(aba, ref):
    """(existe, tem_fórmula, valor_atual_texto)"""
    p, m = _acha(aba, ref)
    if not m: return False, False, None
    corpo = m.group(3) or ""
    if "<f" in corpo: return True, True, None
    at = dict(_attrs(m.group(1))); v = re.search(r"<v>(.*?)</v>", corpo, re.S)
    if at.get("t") == "s" and v: return True, False, s_txt(int(v.group(1)))
    if at.get("t") == "inlineStr": return True, False, "".join(re.findall(r"<t[^>]*>(.*?)</t>", corpo, re.S))
    return True, False, (v.group(1) if v else None)

def _troca(aba, ref, novo_xml):
    p, m = _acha(aba, ref)
    TXT[p] = TXT[p][:m.start()] + novo_xml + TXT[p][m.end():]

def _abre(m, t=None):
    at = [(k, v) for k, v in _attrs(m.group(1)) if k != "t"]
    if t: at.append(("t", t))
    return '<c r="%s"%s' % (m.group(0)[6:m.group(0).index('"', 6)], "".join(' %s="%s"' % kv for kv in at))

def limpar(aba, ref, so_se=None):
    existe, formula, atual = celula_info(aba, ref)
    if not existe: return False
    if formula: LOG.append(("fórmula preservada", aba, ref)); return False
    p, m = _acha(aba, ref)
    if m.group(2) == "/>": return False
    if so_se and not so_se(atual): return False
    _troca(aba, ref, _abre(m) + "/>")
    LOG.append(("limpa", aba, ref, atual)); return True

def gravar(aba, ref, val):
    if val is None or val == "":
        limpar(aba, ref); return
    existe, formula, _ = celula_info(aba, ref)
    if not existe: raise SystemExit("célula %s!%s não existe no arquivo-base (não crio células)" % (aba, ref))
    assert not formula, ("fórmula", aba, ref)
    p, m = _acha(aba, ref)
    if isinstance(val, datetime.datetime):
        d = val - E0; t, v = None, (repr(d.days + d.seconds / 86400) if d.seconds else str(d.days))
    elif isinstance(val, bool):
        t, v = "b", ("1" if val else "0")
    elif isinstance(val, (int, float)):
        t, v = None, repr(val)
    else:
        t, v = "s", str(s_id(str(val)))
    _troca(aba, ref, _abre(m, t) + "><v>%s</v></c>" % v)
    LOG.append(("gravada", aba, ref, val))

# ── 1. abas dos operadores ───────────────────────────────────────────────
if MODO == "setembro":
    # data digitada como "24" (24/01/1900): pelos vizinhos (24/09 dia, 25/09 dia) é 24/09/2026, turno da noite
    gravar("Calibração Blaine", "A51", datetime.datetime(2026, 9, 24))
    # 01/10 (turno 07x19, RENATO) foi lançado na planilha de setembro: passa para a de outubro (linha 6)
    for col in "ACHN": limpar("Calibração Blaine", "%s64" % col)

if MODO in ("padrao", "outubro"):
    for aba, refs in ENTRADAS.items():
        for r in refs:
            if isinstance(r, tuple):        # Blaine J: só apaga o que não for o "---" de fábrica
                limpar(aba, r[0], so_se=lambda v: v is not None and v.strip() != "---")
            else:
                limpar(aba, r)

if MODO == "outubro":
    src = openpyxl.load_workbook(FONTE_OUT, read_only=True, data_only=False)
    setb = openpyxl.load_workbook(BASE, read_only=True, data_only=False)
    def ler(wb, aba, ref):
        ws = wb[aba]; col, lin = CF(ref)
        for row in ws.iter_rows(min_row=lin, max_row=lin, min_col=CI(col), max_col=CI(col)):
            for c in row: return getattr(c, "value", None)
    for aba, refs in OUTUBRO.items():
        for r in refs:
            v = ler(src, aba, r)
            if isinstance(v, str) and v.startswith("="): continue      # fórmula da aba: não mexe
            existe, formula, _ = celula_info(aba, r)
            if formula: continue
            if not existe and v in (None, ""): continue
            gravar(aba, r, v)
    # turno 07x19 de 01/10 que estava na planilha de setembro (linha 64) -> linha 6 de outubro
    for col in "ABCDEFGHLN":          # J (Star) fica com o "---" da linha
        v = ler(setb, "Calibração Blaine", "%s64" % col)
        if isinstance(v, str) and v.startswith("="): continue
        gravar("Calibração Blaine", "%s6" % col, v)

for p, x in TXT.items():
    files[p] = x.encode("utf8")

# ── 2. valores das fórmulas da BD (avalia.py lê as abas dos operadores) ──
def grava_zip(destino):
    with zipfile.ZipFile(destino, "w") as out:
        for n in z.namelist():
            if n not in files: continue
            i = infos[n]
            zi = zipfile.ZipInfo(n, date_time=i.date_time); zi.compress_type = i.compress_type; zi.external_attr = i.external_attr
            out.writestr(zi, files[n])
        for n in files:
            if n not in infos: out.writestr(n, files[n], compress_type=zipfile.ZIP_DEFLATED)

def fecha_sst():
    s2 = re.sub(r'count="\d+" uniqueCount="\d+"', 'count="%d" uniqueCount="%d"' % (
        int(re.search(r'count="(\d+)"', sst).group(1)) + NREFS[0], len(SIS) + len(NOVOS)), sst, 1)
    files["xl/sharedStrings.xml"] = s2.replace("</sst>", "".join(NOVOS) + "</sst>").encode("utf8")

fecha_sst()
tmp = SAIDA + ".tmp.xlsm"; grava_zip(tmp)
cache_f = SAIDA + ".cache.json"
r = subprocess.run([sys.executable, os.path.join(os.path.dirname(os.path.abspath(__file__)), "avalia.py"), tmp, cache_f],
                   capture_output=True, text=True)
print(r.stdout.strip(), r.stderr.strip()[-500:])
assert r.returncode == 0
CACHE = json.load(open(cache_f))["BD_Afericoes"]
os.remove(tmp)
# os textos novos já foram fechados no sst; reabre a lista para a BD
sst = files["xl/sharedStrings.xml"].decode("utf8"); SIS = re.findall(r"<si>.*?</si>|<si/>", sst, re.S); NOVOS = []; NREFS = [0]

# ── 3. BD_Afericoes / BD_Limites ─────────────────────────────────────────
def estilos(p):
    x = files[p].decode("utf8")
    head = re.search(r'<c r="A1" s="(\d+)"', x).group(1)
    data = re.search(r'<c r="C2" s="(\d+)"', x)
    return x, head, data.group(1) if data else None
xa, S_HEAD, S_DATA = estilos(ABA["BD_Afericoes"])
xl_, _, _ = estilos(ABA["BD_Limites"])

def cel(ref, val, st=None):
    s = ' s="%s"' % st if st else ""
    if val is None or val == "": return ""
    if isinstance(val, tuple):
        f = escape(val[1]); cv = CACHE.get(ref)
        if isinstance(cv, (int, float)) and not isinstance(cv, bool):
            return '<c r="%s"%s><f>%s</f><v>%r</v></c>' % (ref, s, f, cv)
        return '<c r="%s"%s t="str"><f>%s</f><v>%s</v></c>' % (ref, s, f, escape(cv if isinstance(cv, str) else ""))
    if isinstance(val, (int, float)): return '<c r="%s"%s><v>%r</v></c>' % (ref, s, val)
    return '<c r="%s"%s t="s"><v>%d</v></c>' % (ref, s, s_id(val))

def troca_dados(x, ultima_col, ultima_lin, linhas_xml, merge=None):
    x = re.sub(r'<dimension ref="[^"]*"/>', '<dimension ref="A1:%s%d"/>' % (ultima_col, ultima_lin), x, 1)
    x = re.sub(r'<pane ([^>]*?)topLeftCell="[^"]*"', r'<pane \1topLeftCell="A2"', x, 1)
    x = re.sub(r'<selection pane="bottomLeft"[^>]*/>', '<selection pane="bottomLeft" activeCell="A2" sqref="A2"/>', x, 1)
    x = re.sub(r"<sheetData>.*</sheetData>|<sheetData/>", "<sheetData>%s</sheetData>" % linhas_xml, x, 1, flags=re.S)
    x = re.sub(r"<mergeCells.*?</mergeCells>", "", x, 1, flags=re.S)
    if merge:
        x = x.replace("</sheetData>", '</sheetData><mergeCells count="1"><mergeCell ref="%s"/></mergeCells>' % merge, 1)
    return x

H = spec.COLS; rows = spec.resolve()
out = ['<row r="1">' + "".join(cel("%s1" % L(i + 1), h, S_HEAD) for i, h in enumerate(H)) + "</row>"]
for n, row in enumerate(rows, start=2):
    out.append('<row r="%d">%s</row>' % (n, "".join(
        cel("%s%d" % (L(i + 1), n), row.get(L(i + 1)), S_DATA if H[i] == "Data" else None) for i in range(len(H)))))
NBD = len(rows) + 1
files[ABA["BD_Afericoes"]] = troca_dados(xa, "Q", NBD, "".join(out)).encode("utf8")

# BD_Limites: parâmetros + o texto "COMO FUNCIONA" que já estava na planilha de setembro
leg_txt = None
m = re.search(r'<c r="A\d+" s="(\d+)" t="s"><v>(\d+)</v></c><c r="B\d+" s="(\d+)"/>', xl_)
if m:
    S_LEG, S_LEG2 = m.group(1), m.group(3); leg_txt = s_txt(int(m.group(2))).rstrip()
HL = ["Chave", "Ensaio", "Parâmetro", "Unidade", "Nominal", "Lim. Inferior", "Lim. Superior", "Tolerância", "Fonte"]
lo = ['<row r="1">' + "".join(cel("%s1" % L(i + 1), h, S_HEAD) for i, h in enumerate(HL)) + "</row>"]
for n, t in enumerate(spec.LIM, start=2):
    lo.append('<row r="%d">%s</row>' % (n, "".join(cel("%s%d" % (L(i + 1), n), v) for i, v in enumerate(t))))
NLIM = len(spec.LIM) + 1
n0 = NLIM + 2
lo.append('<row r="%d">%s</row>' % (n0, cel("A%d" % n0, "COMO FUNCIONA", S_HEAD)))
merge = None
if leg_txt:
    r1, r2 = n0 + 1, n0 + 9
    for r in range(r1, r2 + 1):
        cs = "".join(('<c r="%s%d" s="%s"/>' % (L(c), r, S_LEG2)) if (c, r) != (1, r1) else cel("A%d" % r1, leg_txt, S_LEG) for c in range(1, 10))
        lo.append('<row r="%d">%s</row>' % (r, cs))
    merge = "A%d:I%d" % (r1, r2)
files[ABA["BD_Limites"]] = troca_dados(xl_, "I", (n0 + 9) if leg_txt else n0, "".join(lo), merge).encode("utf8")

# tabelas: só a área muda (mesmos nomes e colunas que o Power BI lê)
for tname, ref in (("tbl_Afericoes", "A1:Q%d" % NBD), ("tbl_Limites", "A1:I%d" % NLIM)):
    for p in [n for n in files if n.startswith("xl/tables/table")]:
        x = files[p].decode("utf8")
        if 'name="%s"' % tname in x:
            x = re.sub(r'( ref=")[A-Z]+\d+:[A-Z]+\d+(")', r'\g<1>%s\2' % ref, x)
            files[p] = x.encode("utf8"); break
    else:
        raise SystemExit("tabela %s não encontrada" % tname)

# ── 4. recalcular ao abrir; calcChain fora (o Excel refaz) ───────────────
wbx = files["xl/workbook.xml"].decode("utf8")
if "fullCalcOnLoad" not in wbx:
    wbx = re.sub(r"<calcPr([^>]*?)/>", r'<calcPr\1 fullCalcOnLoad="1"/>', wbx, 1)
assert 'fullCalcOnLoad="1"' in wbx
files["xl/workbook.xml"] = wbx.encode("utf8")
if "xl/calcChain.xml" in files:
    del files["xl/calcChain.xml"]
    rl = files["xl/_rels/workbook.xml.rels"].decode("utf8")
    files["xl/_rels/workbook.xml.rels"] = re.sub(r'<Relationship [^>]*calcChain[^>]*/>', "", rl).encode("utf8")
    ct = files["[Content_Types].xml"].decode("utf8")
    files["[Content_Types].xml"] = re.sub(r'<Override PartName="/xl/calcChain.xml"[^>]*/>', "", ct).encode("utf8")

fecha_sst()
grava_zip(SAIDA)
json.dump([list(map(str, l)) for l in LOG], open(SAIDA + ".log.json", "w"), ensure_ascii=False, indent=0)
from collections import Counter
print("ok", SAIDA, "| BD linhas:", len(rows), "| operações:", Counter(l[0] for l in LOG))
