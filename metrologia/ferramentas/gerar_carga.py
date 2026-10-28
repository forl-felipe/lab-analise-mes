"""
Gera a carga inicial do sistema Metrologia a partir da planilha de aferição (.xlsm).

Uso:
    python gerar_carga.py "Calibração 10 - Outubro PADRÃO NOVO.xlsm" pasta_de_saida

Saída:
    carga-inicial.xlsx   uma aba por lista, na ordem exata das colunas, pronta
                         para colar no modo de grade do SharePoint
    csv/<Lista>.csv      os mesmos dados em CSV (UTF-8), usados pelo script PnP
    resumo-carga.txt     o que foi gerado e o que precisa de conferência

O script não contém dados do laboratório: tudo é lido da planilha. As posições
das linhas da aba Lançamento seguem o layout "PADRÃO NOVO".
"""

import csv
import datetime as dt
import re
import sys
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from modelo import COLUNAS_AGENDA, lista  # noqa: E402

ENSAIOS_PLANILHA = OrderedDict([
    # nome na agenda da Config : (código, nome no app, NomeBD, formulário)
    ("Calibração Blaine", ("BLA", "Calibração Blaine", "Blaine", "Medição")),
    ("Tambor de abrasão · rotação", ("TAM", "Tambor de abrasão · rotação", "Tambor de Abrasão", "Medição")),
    ("Granulometria fina Alpine", ("ALP", "Granulometria fina Alpine", "Gran. Fina Alpine", "Medição")),
    ("Umidade · analisadores x estufa", ("UMI", "Umidade · analisadores x estufa", "Umidade", "Medição")),
    ("Compressão · prensas", ("COM", "Compressão · prensas", "Compressão", "Medição")),
    ("Granulometria · LTF x LCE", ("GRA", "Granulometria · LTF x LCE", "Granulometria", "Medição")),
    ("Tamboramento 5 kg x 15 kg", ("T515", "Tamboramento 5 kg x 15 kg", "Tamb 5 kg x 15 kg", "Medição")),
    ("Verificação dos peneiradores", ("PEN", "Verificação dos peneiradores", "Verificação Peneiradores", "Checklist")),
    ("Comparativo Fisher", ("FIS", "Comparativo Fisher", "Comparativo Fisher", "Medição")),
])
INSTRUCOES = {
    "BLA": "Informe o resultado de cada aparelho. A diferença para o nominal e a situação são calculadas pelo app.",
    "TAM": "Cronometre o tempo das voltas de cada tambor (minutos e segundos). O app calcula a rotação em rpm.",
    "ALP": "Informe o passante (%) de cada granulômetro.",
    "UMI": "Informe a umidade da estufa (ou as três pesagens) e depois a leitura de cada analisador.",
    "COM": "Informe a resistência média e a velocidade de cada prensa, nas duas faixas.",
    "GRA": "Informe a massa retida (g) em cada fração no LTF e no LCE. Os indicadores são calculados pelo app.",
    "T515": "Informe as massas (g) retidas em +6,3 mm e em -6,3 +0,5 mm para o LTF (5 kg) e o LCE (15 kg).",
    "PEN": "Para cada peneira, marque os cinco itens. Qualquer NÃO OK reprova a peneira. SEM TAG não reprova.",
    "FIS": "Informe a superfície específica medida nos dois aparelhos.",
}
DESCRICAO_PREFIXO = [
    (r"^66PB0?8$", "Blaine manual", "Blaine"),
    (r"^66TA", "Tambor de abrasão", "Tambor"),
    (r"^T03$", "Tambor de abrasão 15 kg (LCE)", "Tambor"),
    (r"^66AG", "Granulômetro Alpine", "Alpine"),
    (r"^66AN", "Analisador de umidade", "Analisador de umidade"),
    (r"^66PS", "Prensa de compressão", "Prensa"),
    (r"^66PN0?9$", "Peneirador LTF", "Peneirador"),
    (r"^66PN0?8$", "Peneirador LCE", "Peneirador"),
    (r"^66PN6", "Peneira do granulômetro Alpine", "Peneira"),
    (r"^66PB", "Fisher (superfície específica)", "Fisher"),
]
SEM_TAG = "P06-10MM"  # peneira de 10 mm do Peneirador 06 sem TAG na planilha (N/CONT)


def tag_peneirador(grupo: str) -> str:
    """'Peneirador NN' -> '<prefixo>PNNN' (padrão das observações da planilha); 'Circulares' -> PN-CIRCULARES (a definir)."""
    m = re.search(r"(\d+)", grupo or "")
    return f"66PN{int(m.group(1)):02d}" if m else "PN-" + re.sub(r"\W+", "", (grupo or "CONJUNTO").upper())


def norm_tag(texto: str) -> str:
    """'PN 614' -> 'PN614'; '66AN11 x Estufa' -> '66AN11'; 'N/CONT' -> P06-10MM."""
    t = (texto or "").strip()
    if t.upper() in ("N/CONT", "N/ CONT"):
        return SEM_TAG
    if t in ("+", "-", ""):
        return ""
    t = t.split(" x ")[0].split(" (")[0]
    return t.replace(" ", "")


def ler_planilha(caminho: Path):
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    lan = {(r, c): v for r, row in enumerate(wb["Lançamento"].iter_rows(max_row=160, max_col=20, values_only=True), 1)
           for c, v in enumerate(row, 1) if v not in (None, "")}
    cfg_rows = list(wb["Config"].iter_rows(max_row=140, max_col=24, values_only=True))
    lim_rows = list(wb["BD_Limites"].iter_rows(values_only=True))
    bd_rows = [r for r in wb["BD_Afericoes"].iter_rows(min_row=2, max_col=21, values_only=True) if r[0]]
    wf = openpyxl.load_workbook(caminho, data_only=False, read_only=True)
    formulas = {(r, c): v for r, row in enumerate(wf["Lançamento"].iter_rows(max_row=160, max_col=20, values_only=True), 1)
                for c, v in enumerate(row, 1) if v not in (None, "")}
    return lan, formulas, cfg_rows, lim_rows, bd_rows


def cel(lan, ref: str):
    m = re.match(r"([A-Z]+)(\d+)", ref)
    col = 0
    for ch in m.group(1):
        col = col * 26 + ord(ch) - 64
    return lan.get((int(m.group(2)), col))


def main(planilha: Path, saida: Path):
    lan, formulas, cfg, lim_rows, bd = ler_planilha(planilha)
    saida.mkdir(parents=True, exist_ok=True)
    avisos: list[str] = []

    # ---------------- Limites e voltas (BD_Limites) ----------------
    limites, voltas = [], {}
    for r in lim_rows[1:]:
        if not r[0] or str(r[0]).startswith("COMO") or len(str(r[0])) > 12:
            continue
        chave = str(r[0]).strip()
        if chave.startswith("V_"):
            voltas[str(r[2]).split()[-1]] = r[4]  # "Voltas 66TA05" -> 66TA05
            continue
        limites.append(OrderedDict(Title=chave, Ensaio=r[1], Parametro=r[2], Unidade=r[3], Nominal=r[4],
                                   LimInferior=r[5], LimSuperior=r[6], Tolerancia=r[7], Fonte=r[8]))
    chaves = {l["Title"] for l in limites}

    # ---------------- Config: agenda, pessoas, escala, parâmetros ----------------
    agenda, pessoas, diam = {}, [], []
    escala_ref, letras07, letras19 = None, [], []
    massa_ltf = massa_lce = None
    for i, row in enumerate(cfg):
        row = list(row) + [None] * 24
        if row[7] in ENSAIOS_PLANILHA:
            agenda[row[7]] = (["X" == str(v).strip().upper() for v in row[8:22]], row[22])
        if row[3] and row[4] in ("A", "B", "C", "D", "ADM") and row[3] != "Nome":
            pessoas.append((row[3].strip(), row[4]))
        if isinstance(row[7], str) and re.match(r"^[+-]\d", row[7]) and isinstance(row[8], (int, float)):
            diam.append((row[7], row[8]))
        if row[7] == "Data de referência (dia 1 do ciclo)":
            escala_ref = row[8]
        if isinstance(row[7], int) and row[8] in ("A", "B", "C", "D") and row[9] in ("A", "B", "C", "D"):
            letras07.append(row[8]); letras19.append(row[9])
        if isinstance(row[1], str) and "massa de ensaio LTF" in row[1]:
            massa_ltf = row[2]
        if isinstance(row[1], str) and "massa de ensaio LCE" in row[1]:
            massa_lce = row[2]

    # ---------------- Ensaios ----------------
    ensaios = []
    for ordem, (nome_ag, (cod, nome, nomebd, form)) in enumerate(ENSAIOS_PLANILHA.items(), 1):
        marcas, freq = agenda.get(nome_ag, ([False] * 14, "Eventual"))
        e = OrderedDict(Title=cod, Nome=nome, NomeBD=nomebd, Formulario=form, Frequencia=freq or "Eventual",
                        Ordem=ordem, Instrucoes=INSTRUCOES[cod], Ativo=True)
        for c, m in zip(COLUNAS_AGENDA, marcas):
            e[c] = m
        ensaios.append(e)

    # ---------------- Pontos de aferição (linhas da aba Lançamento) ----------------
    pontos = []

    def ponto(ens, secao, tag, eqbd, par, un, entrada, chave, fator=None, ra="", rb="", tipo="", ativo=True, conjunto=""):
        if chave and chave not in chaves:
            avisos.append(f"Ponto {ens}/{tag}/{par}: chave de limite {chave} não existe na BD_Limites")
        ordem = sum(1 for p in pontos if p["Ensaio"] == ens) + 1
        codigo = f"{ens}-{ordem:02d}"
        pontos.append(OrderedDict(Title=codigo, Ensaio=ens, Ordem=ordem, Secao=secao, Conjunto=conjunto, TAG=tag, EquipamentoBD=eqbd,
                                  Parametro=par, Unidade=un, Entrada=entrada, ChaveLimite=chave or "", Fator=fator,
                                  RotuloA=ra, RotuloB=rb, TipoRegistro=tipo, Ativo=ativo))

    for r, chave in ((25, "BLAM"), (26, "BLAS"), (27, "BLAA")):
        eq, par = cel(lan, f"C{r}"), cel(lan, f"E{r}")
        tag = norm_tag(eq) if re.match(r"^\d", str(eq)) else ""
        ponto("BLA", "", tag, eq, par, "cm²/g", "Valor", chave, ra="Resultado (cm²/g)", tipo="Calibração", ativo=bool(tag))
    for r in range(33, 38):
        tag = cel(lan, f"C{r}")
        v = voltas.get(tag)
        if v is None:
            avisos.append(f"Tambor {tag}: número de voltas não encontrado na BD_Limites")
        ponto("TAM", "", tag, tag, "Rotação", "rpm", "Tempo", "TAMB", fator=v, ra="Minutos", rb="Segundos", tipo="Calibração")
    for r in range(43, 46):
        tag, pen = cel(lan, f"C{r}"), cel(lan, f"E{r}")
        ponto("ALP", "", tag, tag, f"Passante (peneira {pen})", "%", "Valor", "ALP", ra="Passante (%)", tipo="Calibração")
    for r in range(53, 56):
        tag = cel(lan, f"C{r}")
        ponto("UMI", "", tag, f"{tag} x Estufa", "Umidade", "%", "ComparativoReferencia", "UMI", ra="Analisador (%)", tipo="Comparativo")
    for faixa_row, linhas, chave in ((60, range(62, 67), "COM16"), (67, range(69, 73), "COM12")):
        faixa = cel(lan, f"C{faixa_row}")
        for r in linhas:
            tag = cel(lan, f"C{r}")
            ponto("COM", faixa, tag, tag, f"Resistência {faixa}", "kgf/pel", "Valor", chave, ra="Resistência (kgf/pel)", tipo="Comparativo")
            if cel(formulas, f"M{r}"):  # a planilha só avalia a velocidade onde há fórmula
                ponto("COM", faixa, tag, tag, f"Velocidade {faixa}", "mm/min", "Valor", "VEL", ra="Velocidade (mm/min)", tipo="Verificação")
    dm = dict(diam)
    for r in range(78, 88):
        fr = cel(lan, f"C{r}")
        ponto("GRA", "Massas retidas por fração", "66PN09", "66PN09 (LTF) x 66PN08 (LCE)", fr, "g", "Massa", "",
              fator=dm.get(fr), ra="LTF 66PN09 (g)", rb="LCE 66PN08 (g)", tipo="Comparativo")
    for par in ("+6,3 mm", "-6,3 +0,5 mm"):
        ponto("T515", "Massas retidas", "66TA05", "66TA05 (5 kg) x T03 (15 kg)", par, "g", "Massa", "",
              ra="LTF 5 kg (g)", rb="LCE 15 kg (g)", tipo="Comparativo")
    grupo = ""
    pen_locais = defaultdict(list)
    for r in range(112, 140):
        if cel(lan, f"C{r}") and cel(lan, f"C{r}") not in ("Observação",):
            grupo = cel(lan, f"C{r}")
        malha, tag_txt = cel(lan, f"E{r}"), cel(lan, f"F{r}")
        if not malha or not tag_txt:
            continue
        tag = norm_tag(str(tag_txt))
        pen_locais[tag].append(grupo)
        ponto("PEN", grupo, tag, tag if tag != SEM_TAG else str(tag_txt), f"Malha {malha} mm ({grupo})", "itens OK", "Checklist", "",
              tipo="Verificação", conjunto=tag_peneirador(grupo))
    a, b = cel(lan, "F144") or "66PB01", cel(lan, "H144") or "66PB09"
    a, b = str(a).split(" ")[0], str(b).split(" ")[0]
    ponto("FIS", "", a, f"{a} x {b}", "Superfície específica", "cm²/g", "Comparativo", "FIS", ra=f"{a} (cm²/g)", rb=f"{b} (cm²/g)", tipo="Comparativo")

    for tag, gs in pen_locais.items():
        if len(gs) > 1:
            avisos.append(f"TAG {tag} aparece em mais de um peneirador ({', '.join(gs)}). Confirme qual é a correta antes de imprimir a etiqueta.")

    # ---------------- Equipamentos ----------------
    equip = OrderedDict()

    def add_eq(tag, desc=None, tipo=None, grupo="", ens="", obs="", ativo=True):
        if not tag:
            return
        if tag not in equip:
            d, t = desc, tipo
            if d is None:
                for padrao, dd, tt in DESCRICAO_PREFIXO:
                    if re.search(padrao, tag):
                        d, t = dd, tt
                        break
            equip[tag] = OrderedDict(Title=tag, Descricao=d or tag, TipoEquipamento=t or "", Local="Laboratório físico",
                                     Grupo=grupo, Conjunto="", EnsaioAfericao=ens, ControlaCalibracao=True, PeriodicidadeMeses=12,
                                     UltimaCalibracao=None, ProximaCalibracao=None, ObservacaoCadastro=obs, Ativo=ativo)
        else:
            e = equip[tag]
            if ens and ens not in e["EnsaioAfericao"].split(";"):
                e["EnsaioAfericao"] = ";".join(x for x in (e["EnsaioAfericao"], ens) if x)
            if grupo and grupo not in e["Grupo"]:
                e["Grupo"] = " e ".join(x for x in (e["Grupo"], grupo) if x)
            if obs and obs not in e["ObservacaoCadastro"]:
                e["ObservacaoCadastro"] = "; ".join(x for x in (e["ObservacaoCadastro"], obs) if x)

    for p in pontos:
        if not p["TAG"] or not p["Ativo"]:
            continue
        if p["Ensaio"] == "PEN":
            malha = re.search(r"Malha ([\d,]+) mm", p["Parametro"]).group(1)
            obs = "Sem TAG na planilha (N/CONT): definir TAG e etiquetar" if p["TAG"] == SEM_TAG else ""
            add_eq(p["TAG"], f"Peneira {malha} mm", "Peneira", p["Secao"], "", obs)
            if not equip[p["TAG"]]["Conjunto"]:
                equip[p["TAG"]]["Conjunto"] = p["Conjunto"]
        else:
            add_eq(p["TAG"], ens=p["Ensaio"])
    # Peneiradores: um cadastro por conjunto, com etiqueta própria (o laboratorista verifica pelo QR dele)
    for grupo in dict.fromkeys(p["Secao"] for p in pontos if p["Ensaio"] == "PEN"):
        tp = tag_peneirador(grupo)
        nome = grupo if grupo.lower().startswith("peneirador") else f"Peneiras {grupo.lower()}"
        obs = "" if re.search(r"\d", grupo) else "TAG provisória: definir a TAG oficial deste conjunto"
        if tp in equip:
            equip[tp]["Descricao"] = f"{nome} ({equip[tp]['Descricao']})"
            equip[tp]["TipoEquipamento"] = "Peneirador"
            add_eq(tp, ens="PEN", obs=obs)
        else:
            add_eq(tp, nome, "Peneirador", ens="PEN", obs=obs)
    add_eq("66PN08", ens="")
    add_eq("T03", ens="")
    for p in pontos:
        m = re.search(r"peneira (66PN\d+)", p["Parametro"])
        if m:
            add_eq(m.group(1), ens="")
    for tag, gs in pen_locais.items():
        if len(gs) > 1:
            add_eq(tag, obs=f"TAG repetida na planilha ({' e '.join(gs)}): confirmar")

    # calibração citada nas observações da planilha (ex.: "PN 639 (calibração 06/26)")
    citadas = set()
    for r in bd:
        for m in re.finditer(r"PN\s?0?(\d{3,4})\s*\(\s*calibra[çc][ãa]o\s*(\d{2})/(\d{2})\)", r[15] or "", re.I):
            citadas.add((f"PN{m.group(1)}", m.group(2), m.group(3)))
        for m in re.finditer(r"\(\s*PN\s?(\d{3,4})\s*\)", r[15] or ""):
            if "VENCID" in (r[15] or "").upper():
                citadas.add((f"PN{m.group(1)}", None, None))
        if "VENCID" in (r[15] or "").upper():
            for m in re.finditer(r"PN\s?(\d{3,4})\s+esta vencida", r[15] or "", re.I):
                citadas.add((f"PN{m.group(1)}", None, None))
    mes_ref = next(((mm, aa) for _, mm, aa in citadas if mm), None)
    for tag, mm, aa in sorted(citadas, key=str):
        if tag in equip:
            mm, aa = (mm, aa) if mm else (mes_ref or (None, None))
            if mm:
                ultimo_dia = (dt.date(2000 + int(aa), int(mm) % 12 + 1, 1) - dt.timedelta(days=1)) if int(mm) < 12 else dt.date(2000 + int(aa), 12, 31)
                equip[tag]["ProximaCalibracao"] = ultimo_dia
            equip[tag]["ObservacaoCadastro"] = "; ".join(x for x in (equip[tag]["ObservacaoCadastro"],
                "Vencimento de calibração tirado das observações da planilha: confirmar com a instrumentação") if x)

    # TAGs que só existem no histórico (peneiras trocadas, digitação diferente)
    tags_hist = {norm_tag(r[3]) for r in bd}
    for tag in sorted(t for t in tags_hist if t and t not in equip):
        add_eq(tag, "Peneira (só no histórico)" if tag.startswith("PN") else None,
               "Peneira" if tag.startswith("PN") else None, ens="",
               obs="Aparece só no histórico de outubro: confirmar se ainda existe ou se é erro de digitação", ativo=False)

    # ---------------- Acessos (sem e-mail: a pessoa é escolhida no seletor) ----------------
    acessos = [OrderedDict(Title=n, Pessoa="", Perfil="Laboratorista", Letra=l, Ativo=True,
                           Observacao="Perfil a confirmar (ADM)" if l == "ADM" else "") for n, l in pessoas]

    # ---------------- Configurações ----------------
    configs = [
        ("EscalaDataReferencia", escala_ref.strftime("%Y-%m-%d") if escala_ref else "", "Dia 1 do ciclo da escala de letras (aaaa-mm-dd)"),
        ("EscalaLetras07x19", ",".join(letras07), "Letra de cada dia do ciclo no turno 07x19, separadas por vírgula"),
        ("EscalaLetras19x07", ",".join(letras19), "Letra de cada dia do ciclo no turno 19x07, separadas por vírgula"),
        ("TurnoDiaInicio", "7", "Hora em que começa o turno 07x19"),
        ("TurnoDiaFim", "19", "Hora em que termina o turno 07x19 (começa o 19x07)"),
        ("DiasAvisoCalibracao", "30", "Dias antes do vencimento para mostrar 'vence em breve'"),
        ("T515MassaLTF", str(massa_ltf or 5000), "Massa de ensaio do LTF no tamboramento (g)"),
        ("T515MassaLCE", str(massa_lce or 15000), "Massa de ensaio do LCE no tamboramento (g)"),
        ("InicioControle", escala_ref.strftime("%Y-%m-%d") if escala_ref else "", "Data a partir da qual os ensaios previstos contam no Power BI"),
        ("EmailsRelatorioDiario", "", "E-mails que recebem o relatório diário, separados por ponto e vírgula"),
        ("EmailsAlertaNaoConforme", "", "E-mails extras para alerta de não conformidade (a instrumentação já recebe)"),
        ("LinkApp", "", "Link da Web do app (Power Apps > Detalhes)"),
        ("PastaImagensQR", "", "Endereço da biblioteca ConsultaQR (ex.: https://EMPRESA.sharepoint.com/sites/SITE/ConsultaQR). O QR Code de cada etiqueta é este endereço + /TAG.png"),
        ("EmailsAvisoImagem", "", "Quem recebe o aviso para atualizar a imagem de consulta (calibração registrada e calibração vencida), separados por ponto e vírgula"),
    ]
    configs = [OrderedDict(Title=k, Valor=v, Descricao=d) for k, v, d in configs]
    if len(letras07) != 8:
        avisos.append("Escala de letras: esperado ciclo de 8 dias")

    # ---------------- Histórico: Afericoes e Lancamentos ----------------
    def itens_nao_ok(obs):
        itens = ["Condições das telas", "Inclinação das telas", "Vibração do peneirador", "Travamento das peneiras", "Calibração"]
        return "; ".join(i for i in itens if i.lower() in (obs or "").lower())

    # Na planilha um mesmo código LCP pode juntar mais de um ensaio; no sistema novo cada
    # lançamento é de um ensaio só, então esses códigos ganham o sufixo do ensaio.
    cod_por_nomebd = {v[2]: v[0] for v in ENSAIOS_PLANILHA.values()}
    ensaios_do_lote = defaultdict(set)
    for r in bd:
        ensaios_do_lote[(r[16] or "").split(" | ")[0].strip()].add(r[0])
    afer, lotes = [], OrderedDict()
    for r in bd:
        lote = (r[16] or "").split(" | ")[0].strip()
        if len(ensaios_do_lote[lote]) > 1:
            lote = f"{lote}-{cod_por_nomebd.get(r[0], 'X')}"
        rec = OrderedDict(
            Title=lote, Ensaio=r[0], TipoRegistro=r[1], DataEnsaio=r[2].date() if r[2] else None, Equipamento=r[3],
            TAG=norm_tag(r[3]), Parametro=r[4], Unidade=r[5], Valor=r[6], Referencia=r[7], Diferenca=r[8],
            LimInferior=r[9], LimSuperior=r[10], Tolerancia=r[11], Resultado=r[12], Responsavel=r[13],
            ResponsavelEmail="", Letra=r[14], Turno=r[17], Observacao=(r[15] or "").strip(), Causa="",
            ItensNaoOK=itens_nao_ok(r[15]) if r[0] == "Verificação Peneiradores" and r[12] == "Não conforme" else "",
            RegistradoEm=r[18], TurnoRegistro=r[19], ForaDoTurno=r[20], Origem="Planilha")
        afer.append(rec)
        L = lotes.setdefault(lote, OrderedDict(Title=lote, Ensaio=r[0], DataEnsaio=rec["DataEnsaio"], Turno=r[17], Letra=r[14],
                                               Responsavel=r[13], ResponsavelEmail="", QtdResultados=0, QtdNaoConformes=0,
                                               TAGsNaoConformes="", Causa="", Observacao="", RegistradoEm=r[18],
                                               ForaDoTurno=r[20], Situacao="Ativo", Origem="Planilha"))
        L["QtdResultados"] += 1
        if r[12] == "Não conforme":
            L["QtdNaoConformes"] += 1
            L["TAGsNaoConformes"] = ", ".join(x for x in (L["TAGsNaoConformes"], rec["TAG"]) if x)
            L["Observacao"] = rec["Observacao"]

    # ---------------- Escrita ----------------
    tabelas = OrderedDict([
        ("Acessos", acessos), ("Equipamentos", list(equip.values())), ("Ensaios", ensaios),
        ("PontosAfericao", pontos), ("Limites", limites), ("Configuracoes", configs),
        ("Afericoes", afer), ("Lancamentos", list(lotes.values())),
    ])
    escrever_excel(saida / "carga-inicial.xlsx", tabelas, avisos)
    (saida / "csv").mkdir(exist_ok=True)
    for nome, linhas in tabelas.items():
        escrever_csv(saida / "csv" / f"{nome}.csv", linhas)

    resumo = [f"Carga gerada em {dt.datetime.now():%d/%m/%Y %H:%M} a partir de {planilha.name}", ""]
    resumo += [f"{n:<16} {len(v):>4} linhas" for n, v in tabelas.items()]
    resumo += ["", "PONTOS PARA CONFERIR:"] + [f"- {a}" for a in avisos] if avisos else []
    (saida / "resumo-carga.txt").write_text("\n".join(resumo) + "\n", encoding="utf-8")
    escrever_listas_etiquetas(saida, list(equip.values()))
    print("\n".join(resumo))


def situacao_imagem(e, hoje=None):
    """Situação mostrada na imagem de consulta: NÃO UTILIZAR só quando a calibração está vencida."""
    hoje = hoje or dt.date.today()
    if not e["ControlaCalibracao"]:
        return "SEM CALIBRAÇÃO PERIÓDICA"
    if not e["ProximaCalibracao"]:
        return "A CONFIRMAR"
    return "NÃO UTILIZAR" if e["ProximaCalibracao"] < hoje else "CALIBRADO"


def escrever_listas_etiquetas(saida: Path, equipamentos):
    """Listas prontas para colar nos geradores (etiquetas físicas e imagens de consulta)."""
    ativos = [e for e in equipamentos if e["Ativo"]]
    eh_conjunto = lambda e: e["TipoEquipamento"] == "Peneirador" and "PEN" in e["EnsaioAfericao"].split(";")
    linha = lambda e: f"{e['Title']}; {e['Descricao']}; {e['Grupo'] if e['TipoEquipamento'] == 'Peneira' and e['Grupo'] else e['Local']}"
    grupos = {
        "etiquetas-equipamentos.txt": [e for e in ativos if e["TipoEquipamento"] != "Peneira" and not eh_conjunto(e)],
        "etiquetas-peneiradores.txt": [e for e in ativos if eh_conjunto(e)],
        "etiquetas-peneiras.txt": [e for e in ativos if e["TipoEquipamento"] == "Peneira"],
    }
    for nome, lista in grupos.items():
        (saida / nome).write_text("\n".join(linha(e) for e in lista) + "\n", encoding="utf-8")
    fmt = lambda d: d.strftime("%d/%m/%Y") if d else ""
    img = [f"{e['Title']}; {e['Descricao']}; {e['Local']}; {situacao_imagem(e)}; {fmt(e['UltimaCalibracao'])}; "
           f"{fmt(e['ProximaCalibracao'])}; ; {e['PeriodicidadeMeses'] if e['ControlaCalibracao'] else ''}" for e in ativos]
    (saida / "imagens-consulta.txt").write_text("\n".join(img) + "\n", encoding="utf-8")
    print(f"listas: {', '.join(f'{n} ({len(l)})' for n, l in grupos.items())}, imagens-consulta.txt ({len(img)})")


def valor_csv(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "1" if v else "0"
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%dT%H:%M:%S")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return repr(round(v, 10))
    return str(v)


def escrever_csv(caminho: Path, linhas):
    if not linhas:
        return
    with caminho.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(list(linhas[0].keys()))
        for l in linhas:
            w.writerow([valor_csv(v) for v in l.values()])


def escrever_excel(caminho: Path, tabelas, avisos):
    wb = openpyxl.Workbook()
    capa = wb.active
    capa.title = "LEIA-ME"
    azul = PatternFill("solid", fgColor="0B2540")
    texto = [
        ("CARGA INICIAL · SISTEMA METROLOGIA", True),
        ("", False),
        ("Cada aba corresponde a uma lista do SharePoint, com as colunas na mesma ordem.", False),
        ("Como carregar: abra a lista no SharePoint > Editar no modo de grade > selecione a 1ª célula vazia > cole (Ctrl+V).", False),
        ("Copie a partir da linha 2 (sem o cabeçalho). Ordem recomendada: Configuracoes, Limites, Ensaios, PontosAfericao, Equipamentos, Afericoes, Lancamentos, Acessos.", False),
        ("Colunas Sim/Não aparecem como VERDADEIRO/FALSO: o modo de grade aceita assim.", False),
        ("Acessos: a coluna Pessoa precisa ser preenchida no SharePoint escolhendo a pessoa (seletor). Use esta aba como lista de conferência.", False),
        ("Equipamentos: datas de calibração estão vazias, exceto as citadas nas observações da planilha. Preencha com a instrumentação.", False),
        ("", False),
        ("PONTOS PARA CONFERIR", True),
    ] + [(f"• {a}", False) for a in avisos]
    for i, (t, b) in enumerate(texto, 1):
        c = capa.cell(i, 1, t)
        c.font = Font(bold=b, size=14 if i == 1 else 11, color="0B2540" if b else "1F2937")
    capa.column_dimensions["A"].width = 150

    for nome, linhas in tabelas.items():
        ws = wb.create_sheet(nome)
        if not linhas:
            continue
        cab = list(linhas[0].keys())
        L = lista(nome)
        for j, h in enumerate(cab, 1):
            c = ws.cell(1, j, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = azul
            c.alignment = Alignment(horizontal="center")
            desc = "Título: " + L.titulo_significa if h == "Title" else next((x.descricao for x in L.colunas if x.nome == h), "")
            if desc:
                from openpyxl.comments import Comment
                c.comment = Comment(desc, "Metrologia")
            ws.column_dimensions[get_column_letter(j)].width = max(10, min(45, len(h) + 4))
        for i, l in enumerate(linhas, 2):
            for j, v in enumerate(l.values(), 1):
                c = ws.cell(i, j, v)
                if isinstance(v, dt.datetime):
                    c.number_format = "DD/MM/YYYY HH:MM"
                elif isinstance(v, dt.date):
                    c.number_format = "DD/MM/YYYY"
        ws.freeze_panes = "A2"
    wb.save(caminho)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    main(Path(sys.argv[1]), Path(sys.argv[2]))
